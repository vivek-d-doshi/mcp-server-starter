"""PDF, DOCX, XLSX. Parsed from memory; never executed, never written. Guards against zip bombs and huge documents."""

from __future__ import annotations

import io
import zipfile

from filelens.extract import ExtractError, Extraction

MAX_UNZIPPED = 200 * 1024 * 1024
MAX_RATIO = 200


def _check_zip(data: bytes) -> None:
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile:
        raise ExtractError("The file is not a valid Office document (corrupt zip).") from None
    total = sum(i.file_size for i in infos)
    if total > MAX_UNZIPPED or (len(data) and total / len(data) > MAX_RATIO and total > 10 * 1024 * 1024):
        raise ExtractError("Refused: the document expands to an unsafe size (possible zip bomb).")


def extract_pdf(data: bytes, *, max_pages: int, max_chars: int) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("The PDF is password-protected. FileLens does not ask for or guess passwords.")
        total = len(reader.pages)
    except ExtractError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise ExtractError(f"The PDF could not be read ({type(exc).__name__}).") from None
    parts: list[str] = []
    empty: list[int] = []
    used = 0
    last = 0
    for n in range(min(total, max_pages)):
        try:
            page_text = (reader.pages[n].extract_text() or "").strip()
        except Exception:  # a single bad page must not sink the document
            page_text = ""
        if not page_text:
            empty.append(n + 1)
        block = f"[page {n + 1}]\n{page_text}"
        if used + len(block) > max_chars and parts:
            break
        parts.append(block)
        used += len(block)
        last = n + 1
    ex = Extraction("\n\n".join(parts), "pypdf", "page", total, (1, last) if last else None, truncated=last < total)
    info = reader.metadata
    if info:
        for key, label in (("/Title", "title"), ("/Author", "author")):
            if info.get(key):
                ex.meta[label] = str(info[key])[:200]
    if empty:
        shown = ", ".join(map(str, empty[:10]))
        ex.warnings.append(f"No extractable text on page(s) {shown}. They may be scanned images; OCR is not supported.")
    if last < total:
        ex.continue_hint = f"Only pages 1-{last} of {total} were included (page/size limit)."
    return ex


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    esc = [[c.replace("|", "\\|").replace("\n", " ") for c in r] for r in rows]
    lines = ["| " + " | ".join(esc[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in esc[1:]]
    return "\n".join(lines)


def extract_docx(data: bytes, *, max_chars: int) -> Extraction:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError(f"The Word document could not be read ({type(exc).__name__}).") from None
    out: list[str] = []
    headings = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.startswith("Heading"):
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                out.append("#" * min(level, 6) + " " + text)
                headings += 1
            elif style.startswith("List"):
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            out.append(_md_table([[c.text.strip() for c in row.cells] for row in table.rows]))
    joined = "\n\n".join(out)
    ex = Extraction(joined[:max_chars], "python-docx", "section", headings or None, None, truncated=len(joined) > max_chars)
    props = document.core_properties
    for label, value in (("title", props.title), ("author", props.author)):
        if value:
            ex.meta[label] = value[:200]
    if ex.truncated:
        ex.continue_hint = "Document text was cut at the output limit."
    return ex


def extract_xlsx(data: bytes, *, max_rows: int, max_chars: int) -> Extraction:
    import openpyxl

    _check_zip(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError(f"The spreadsheet could not be read ({type(exc).__name__}).") from None
    parts: list[str] = []
    sheet_names = wb.sheetnames
    for name in sheet_names[:10]:
        ws = wb[name]
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows + 1:
                break
            rows.append(["" if c is None else str(c) for c in row[:20]])
        dims = f"{ws.max_row or '?'} rows x {ws.max_column or '?'} columns"
        parts.append(f"[sheet {name} · {dims}]\n{_md_table(rows) if rows else '(empty)'}")
    joined = "\n\n".join(parts)
    ex = Extraction(
        joined[:max_chars],
        "openpyxl",
        "sheet",
        len(sheet_names),
        (1, min(len(sheet_names), 10)),
        truncated=len(joined) > max_chars,
    )
    ex.meta["sheets"] = ", ".join(sheet_names[:20])
    ex.warnings.append(
        f"Showing the first {max_rows} rows and 20 columns per sheet. Formulas are shown as their last saved values (empty if never calculated)."
    )
    return ex
