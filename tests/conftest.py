import os
from pathlib import Path

import pytest

from filelens.config import build_settings
from filelens.sandbox import Sandbox


@pytest.fixture
def tree(tmp_path):
    """root/ is allowed. outside/ is not. outside/secret.txt must never be readable."""
    root = tmp_path / "root"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    (root / "hello.txt").write_text("hello\nworld\n")
    (root / "sub").mkdir()
    (root / "sub" / "deep.md").write_text("# Deep\n")
    (outside / "secret.txt").write_text("TOP SECRET")
    return Path(os.path.realpath(tmp_path)), root, outside


@pytest.fixture
def sandbox(tree):
    _, root, _ = tree
    return Sandbox(build_settings([str(root)]))


# ---- generated documents (real files, not mocks) -----------------------------------------------------------------
def make_pdf(pages, password=None):
    import io

    from pypdf import PdfWriter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        if text:
            c.drawString(72, 750, text)
        c.showPage()
    c.save()
    if not password:
        return buf.getvalue()
    writer = PdfWriter(clone_from=io.BytesIO(buf.getvalue()))
    writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx():
    import io

    import docx

    d = docx.Document()
    d.core_properties.title = "Quarterly Notes"
    d.add_heading("Title", level=1)
    d.add_paragraph("Intro paragraph.")
    d.add_paragraph("bullet one", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Name", "Qty", "Pen", "3"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Price"])
    ws.append(["Pen", 1.5])
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
